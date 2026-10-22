"""Decks that exercise what PPTX import has to draw: every built-in table style, and charts.

    python3 scripts/pptx-fidelity/corpus.py <out dir>

Writes `tables.pptx` (each of PowerPoint's 74 built-in table styles named by id only, as files
from other tools do, plus borders, merges, diagonals and vertical text set on the cells) and
`charts.pptx` (one chart per slide: every chart type python-pptx writes, then hand-written parts
for the types it cannot, such as stock, bar of pie, surface, combinations on two axes and the
Office 2016 charts). Feed them to `compare.py --import`. Needs python-pptx; not part of `pnpm test`.
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import json
import re

ROOT = Path(__file__).resolve().parents[2]
# the 74 built-in table styles, by id and name, as the importer carries them
STYLES = [tuple(json.loads(m)[:2]) for m in re.findall(r'^  (\[".*\]),$', (ROOT / 'packages/core/src/pptxTableStyles.data.ts').read_text(), re.M)]


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p


def title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.6))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)


def table(slide, style_id, flags, rows=5, cols=4, top=1.1):
    shape = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(top), Inches(11.7), Inches(0.5 * rows))
    t = shape.table
    for r in range(rows):
        for c in range(cols):
            t.cell(r, c).text = ['Region', 'Q1', 'Q2', 'Q3'][c] if r == 0 else (['North', 'South', 'East', 'West', 'Total', 'Other', 'More'][r - 1] if c == 0 else str(10 * r + c))
            for p in t.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(16)
    pr = shape._element.graphic.graphicData.tbl.tblPr
    for k in ('firstRow', 'bandRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol'):
        if k in pr.attrib:
            del pr.attrib[k]
    for k in flags:
        pr.set(k, '1')
    sid = pr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(pr, qn('a:tableStyleId'))
    sid.text = style_id
    return t


def tables(out):
    p = deck()
    blank = p.slide_layouts[6]
    for i in range(0, len(STYLES), 2):
        s = p.slides.add_slide(blank)
        names = []
        for j, (sid, name) in enumerate(STYLES[i:i + 2]):
            flags = ['firstRow', 'bandRow'] if j == 0 else ['firstRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol']
            table(s, sid, flags, rows=5, top=1.1 + j * 3.1)
            names.append(name)
        title(s, ' / '.join(names))
    # cell formatting over the style: own fills and borders, merges, diagonals, vertical text
    s = p.slides.add_slide(blank)
    title(s, 'Cell borders, merges, diagonals, vertical text')
    t = table(s, '{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}', ['firstRow', 'bandRow'], rows=5)
    t.cell(1, 1).merge(t.cell(2, 2))
    t.cell(1, 1).text = 'Merged 2×2'
    t.cell(3, 3).fill.solid()
    t.cell(3, 3).fill.fore_color.rgb = RGBColor(0xFF, 0xC0, 0x00)
    pr = t.cell(4, 1)._tc.get_or_add_tcPr()
    for tag, w, colour, cmpd in (('a:lnL', 38100, 'C00000', 'sng'), ('a:lnR', 38100, 'C00000', 'sng'), ('a:lnT', 38100, 'C00000', 'sng'), ('a:lnB', 57150, '1F4E79', 'dbl'), ('a:lnTlToBr', 12700, '000000', 'sng')):
        ln = etree.SubElement(pr, qn(tag), w=str(w), cmpd=cmpd)
        etree.SubElement(etree.SubElement(ln, qn('a:solidFill')), qn('a:srgbClr'), val=colour)
    v = t.cell(4, 3)._tc.get_or_add_tcPr()
    v.set('vert', 'vert270')
    t.cell(4, 3).text = 'Up'
    p.save(out / 'tables.pptx')


if __name__ == '__main__':
    out = Path(sys.argv[1] if len(sys.argv) > 1 else '.')
    out.mkdir(parents=True, exist_ok=True)
    tables(out)
    sys.path.insert(0, str(Path(__file__).parent))
    from corpus_charts import charts
    charts(out)
